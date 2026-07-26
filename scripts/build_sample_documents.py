from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import letter
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"
SAMPLES.mkdir(exist_ok=True)

TITLE = "Synthetic Sick Leave Procedure"
SUBTITLE = "Northstar Café Group | Approved demonstration policy | Version 1"
STEPS = [
    "Contact your manager before your scheduled shift using the approved team contact method.",
    "State that you cannot attend work and, if known, provide your expected return date.",
    "Do not send detailed medical information through routine operational channels.",
    "If the absence continues beyond the initial notification, contact HR for the next approved step.",
]

def set_font(run, size=11, color="16221D", bold=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold

def add_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "174F3F")
    borders.append(bottom)
    p_pr.append(borders)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Aptos", Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for style_name, size, color in [("Heading 1", 16, "174F3F"), ("Heading 2", 13, "174F3F")]:
    style = doc.styles[style_name]
    style.font.name, style.font.size, style.font.color.rgb = "Aptos Display", Pt(size), RGBColor.from_string(color)
    style.font.bold = True

kicker = doc.add_paragraph()
kicker.paragraph_format.space_after = Pt(8)
set_font(kicker.add_run("SYNTHETIC POLICY • EMPLOYEE OPERATIONS"), 9, "174F3F", True)
title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(5)
set_font(title.add_run(TITLE), 24, "16221D", True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
set_font(subtitle.add_run(SUBTITLE), 11, "637169")
add_bottom_rule(subtitle)

notice = doc.add_paragraph()
notice.paragraph_format.space_before = Pt(10)
notice.paragraph_format.space_after = Pt(14)
notice.paragraph_format.left_indent = Inches(.18)
set_font(notice.add_run("DEMONSTRATION ONLY. "), 10, "8A5A00", True)
set_font(notice.add_run("All organizations, people, and records associated with this policy are fictional. Do not use this document as legal or employment advice."), 10, "5F4A14")

doc.add_heading("Purpose", level=1)
doc.add_paragraph("This fictional procedure explains the operational notification steps an employee follows when illness prevents them from attending a scheduled shift.")
doc.add_heading("What to do", level=1)
for step in STEPS:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    p.add_run(step)
doc.add_heading("Manager response", level=1)
doc.add_paragraph("The manager records that the employee reported an absence, confirms the next communication point, and escalates questions about extended absence or accommodations to HR. Managers must not request unnecessary medical details.")
doc.add_heading("Limits and escalation", level=1)
doc.add_paragraph("This procedure does not approve leave, determine eligibility, make a medical assessment, or replace applicable employment standards. When the approved policy does not answer a question, the employee should contact their manager or HR.")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Synthetic data only | Northstar Café Group | Version 1"), 8, "637169")
doc.core_properties.title = TITLE
doc.core_properties.author = "Private Company AI synthetic fixture"
doc.save(SAMPLES / "synthetic-sick-leave-policy.docx")

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="Kicker", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=HexColor("#174F3F"), spaceAfter=8))
styles.add(ParagraphStyle(name="PolicyTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=28, textColor=HexColor("#16221D"), alignment=TA_LEFT, spaceAfter=5))
styles.add(ParagraphStyle(name="PolicySub", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=13, textColor=HexColor("#637169"), spaceAfter=16))
styles.add(ParagraphStyle(name="PolicyH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=HexColor("#174F3F"), spaceBefore=12, spaceAfter=7))
styles.add(ParagraphStyle(name="PolicyBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15, textColor=HexColor("#16221D"), spaceAfter=7))

def footer_pdf(canvas, document):
    canvas.saveState(); canvas.setFont("Helvetica", 8); canvas.setFillColor(HexColor("#637169")); canvas.drawCentredString(letter[0]/2, 35, f"Synthetic data only | Version 1 | Page {document.page}"); canvas.restoreState()

story = [Paragraph("SYNTHETIC POLICY • EMPLOYEE OPERATIONS", styles["Kicker"]), Paragraph(TITLE, styles["PolicyTitle"]), Paragraph(SUBTITLE, styles["PolicySub"]), Paragraph("<b>DEMONSTRATION ONLY.</b> All associated organizations, people, and records are fictional. This is not legal or employment advice.", styles["PolicyBody"]), Paragraph("Purpose", styles["PolicyH1"]), Paragraph("This fictional procedure explains the operational notification steps an employee follows when illness prevents them from attending a scheduled shift.", styles["PolicyBody"]), Paragraph("What to do", styles["PolicyH1"]), ListFlowable([ListItem(Paragraph(x, styles["PolicyBody"]), leftIndent=12) for x in STEPS], bulletType="1", start="1", leftIndent=22, bulletFontName="Helvetica-Bold"), Paragraph("Manager response", styles["PolicyH1"]), Paragraph("The manager records that the employee reported an absence, confirms the next communication point, and escalates questions about extended absence or accommodations to HR. Managers must not request unnecessary medical details.", styles["PolicyBody"]), Paragraph("Limits and escalation", styles["PolicyH1"]), Paragraph("This procedure does not approve leave, determine eligibility, make a medical assessment, or replace applicable employment standards. When the approved policy does not answer a question, the employee should contact their manager or HR.", styles["PolicyBody"]) ]
pdf = SimpleDocTemplate(str(SAMPLES / "synthetic-sick-leave-policy.pdf"), pagesize=letter, rightMargin=72, leftMargin=72, topMargin=68, bottomMargin=58, title=TITLE, author="Private Company AI synthetic fixture")
pdf.build(story, onFirstPage=footer_pdf, onLaterPages=footer_pdf)
